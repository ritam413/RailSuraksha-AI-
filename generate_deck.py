import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_complete_sih_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Clean Mintlify / Indian Railways Palette
    c_bg = RGBColor(240, 246, 252)        # #F0F6FC Base
    c_card = RGBColor(255, 255, 255)      # #FFFFFF Card
    c_border = RGBColor(208, 223, 238)    # #D0DFEE Border
    c_blue = RGBColor(43, 127, 255)       # #2B7FFF Signal Blue
    c_navy = RGBColor(15, 23, 42)         # #0F172A Ink Slate
    c_slate = RGBColor(71, 85, 105)       # #475569 Slate
    c_light_blue = RGBColor(230, 240, 250) # #E6F0FA
    c_green = RGBColor(16, 185, 129)      # #10B981 Green
    c_red = RGBColor(239, 68, 68)         # #EF4444 Red
    c_dark_blue = RGBColor(30, 58, 138)   # #1E3A8A Dark Blue

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="SMART INDIA HACKATHON | MINISTRY OF RAILWAYS"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_blue

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = c_navy

    # =========================================================================
    # SLIDE 1: Title & Overview
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)

    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.9), Inches(11.333), Inches(5.7))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_card
    card1.line.color.rgb = c_border
    card1.line.width = Pt(1.5)

    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = tf1.margin_top = Inches(0.8)

    p = tf1.paragraphs[0]
    p.text = "SMART INDIA HACKATHON | HARDWARE / SOFTWARE EDITION"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = c_blue

    p2 = tf1.add_paragraph()
    p2.text = "Automatic Block Planning & Corridor Optimization System"
    p2.font.size = Pt(27)
    p2.font.bold = True
    p2.font.color.rgb = c_navy
    p2.space_before = Pt(14)
    p2.space_after = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Integrated Multi-Department Maintenance Scheduling across TMS, SMMS, TDMS, and COA Train Traffic Data"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = c_dark_blue
    p3.space_after = Pt(20)

    p4 = tf1.add_paragraph()
    p4.text = "Organization: Ministry of Railways  |  Target: Civil Engineering, TRD (Traction), S&T, and Operations (COA)"
    p4.font.size = Pt(11.5)
    p4.font.color.rgb = c_slate

    # =========================================================================
    # SLIDE 2: Problem Understanding & Existing Operational Gaps
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "Problem Understanding: Current Decentralized Block Planning")

    # 3 Column Cards for the 3 Fragmented Systems
    systems = [
        ("TMS (Track)", "Civil Engineering", ["Maintains rail defects, fractures, and overdue track tamping data in isolation.", "Requests track blocks independently via BDMS without visibility into other departments."], c_blue),
        ("TDMS (Traction)", "Electrical / TRD", ["Manages OHE catenary wear, insulator washing, and power disconnections.", "Requires power cut blocks that often clash or occur days apart from track work."], c_navy),
        ("SMMS (Signaling)", "Signal & Telecom", ["Tracks point machine overhauls, signal aspects, and track circuit calibrations.", "Manual block coordination causes repeated train speed restrictions."], c_dark_blue)
    ]

    for i, (title_s, sub_s, items_s, accent_s) in enumerate(systems):
        cx = Inches(0.8 + i * 4.04)
        c_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.5), Inches(3.64), Inches(3.6))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = c_card
        c_shape.line.color.rgb = c_border
        c_shape.line.width = Pt(1.5)

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)

        p_h = tf.paragraphs[0]
        p_h.text = sub_s.upper()
        p_h.font.size = Pt(9)
        p_h.font.bold = True
        p_h.font.color.rgb = accent_s

        p_t = tf.add_paragraph()
        p_t.text = title_s
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = c_navy
        p_t.space_before = Pt(4)
        p_t.space_after = Pt(10)

        for it in items_s:
            pi = tf.add_paragraph()
            pi.text = "• " + it
            pi.font.size = Pt(10)
            pi.font.color.rgb = c_slate
            pi.space_after = Pt(8)

    # Bottom Callout Box for COA & Result
    b_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.6))
    b_card.fill.solid()
    b_card.fill.fore_color.rgb = c_light_blue
    b_card.line.color.rgb = c_blue
    b_card.line.width = Pt(1)

    btf = b_card.text_frame
    btf.word_wrap = True
    btf.margin_left = btf.margin_right = btf.margin_top = Inches(0.3)

    bp1 = btf.paragraphs[0]
    bp1.text = "THE BOTTLENECK IN CURRENT BDMS (BLOCK DEMAND MANAGEMENT SYSTEM)"
    bp1.font.size = Pt(10)
    bp1.font.bold = True
    bp1.font.color.rgb = c_blue

    bp2 = btf.add_paragraph()
    bp2.text = "Without automated integration with the Control Office Application (COA) train timetable and goods forecast, maintenance blocks are approved manually. This results in severe corridor downtime, frequent speed restrictions, and disrupted freight operations."
    bp2.font.size = Pt(11)
    bp2.font.color.rgb = c_navy
    bp2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 3: Proposed Solution (Pictorial Architecture)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "Proposed Solution: Unified Block Planning Architecture")

    cards_data = [
        {
            "num": "01",
            "title": "Unified Ingestion",
            "subtitle": "Breaks Departmental Silos",
            "items": [
                "• TMS (Track): Geometry defects, rail fractures, and overdue tamping.",
                "• SMMS (Signals): Point machines, relays, and interlocking overhauls.",
                "• TDMS (Traction/OHE): Power disconnections & catenary maintenance.",
                "• COA (Traffic): Passenger timetables & goods freight density forecasts."
            ],
            "accent": c_blue
        },
        {
            "num": "02",
            "title": "Optimization Engine",
            "subtitle": "Constraint-Aware Solver",
            "items": [
                "• Dynamic Traffic Gaps: Detects natural off-peak freight gaps in COA.",
                "• Automated Shadow Blocking: Merges compatible multi-department tasks.",
                "• Urgency & Risk Scoring: Prioritizes P1 safety over routine tasks.",
                "• Multi-Horizon Output: Auto-generates daily, weekly, & monthly plans."
            ],
            "accent": c_dark_blue
        },
        {
            "num": "03",
            "title": "Controller Cockpit",
            "subtitle": "Human-in-the-Loop Execution",
            "items": [
                "• Corridor String Chart: Live visual map of train paths vs block windows.",
                "• 1-Click Controller Approval: Fast review, modify, or override gates.",
                "• Explainable AI Dossier: 4-step audit trail showing why slots were chosen.",
                "• Real-Time Gang Tracking: Live field crew safety & clearance timers."
            ],
            "accent": c_green
        }
    ]

    for i, data in enumerate(cards_data):
        cx = Inches(0.8 + i * 4.04)
        c_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.5), Inches(3.64), Inches(3.8))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = c_card
        c_shape.line.color.rgb = c_border
        c_shape.line.width = Pt(1.5)

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)

        p_num = tf.paragraphs[0]
        p_num.text = f"STAGE {data['num']}  |  {data['subtitle'].upper()}"
        p_num.font.size = Pt(9)
        p_num.font.bold = True
        p_num.font.color.rgb = data["accent"]

        p_t = tf.add_paragraph()
        p_t.text = data["title"]
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = c_navy
        p_t.space_before = Pt(4)
        p_t.space_after = Pt(10)

        for item in data["items"]:
            pi = tf.add_paragraph()
            pi.text = item
            pi.font.size = Pt(10.5)
            pi.font.color.rgb = c_slate
            pi.space_after = Pt(6)

    # Bottom Metrics Banner
    m_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.5), Inches(11.733), Inches(1.4))
    m_shape.fill.solid()
    m_shape.fill.fore_color.rgb = c_light_blue
    m_shape.line.color.rgb = c_blue
    m_shape.line.width = Pt(1)

    metrics = [
        ("35% - 40%", "Corridor Downtime Reduction", "Via Multi-Dept Joint Shadow Blocking"),
        ("0", "Passenger Path Disruptions", "Optimized during off-peak traffic gaps"),
        ("< 30s", "Automated Schedule Solver", "Replaces multi-day manual coordination"),
        ("100%", "Explainable Audit Trail", "Tamper-evident RDSO compliance dossier")
    ]

    for j, (val, title_m, desc_m) in enumerate(metrics):
        mx = Inches(0.9 + j * 2.9)
        mbox = s3.shapes.add_textbox(mx, Inches(5.55), Inches(2.75), Inches(1.3))
        mtf = mbox.text_frame
        mtf.word_wrap = True
        mtf.margin_left = mtf.margin_top = mtf.margin_right = mtf.margin_bottom = 0
        
        mp1 = mtf.paragraphs[0]
        mp1.text = val
        mp1.font.size = Pt(20)
        mp1.font.bold = True
        mp1.font.color.rgb = c_blue

        mp2 = mtf.add_paragraph()
        mp2.text = title_m
        mp2.font.size = Pt(10)
        mp2.font.bold = True
        mp2.font.color.rgb = c_navy

        mp3 = mtf.add_paragraph()
        mp3.text = desc_m
        mp3.font.size = Pt(8.5)
        mp3.font.color.rgb = c_slate

    # =========================================================================
    # SLIDE 4: Core Innovation — Automated Shadow Blocking
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "Core Technical Innovation: Multi-Department Shadow Blocking")

    # Left: Fragmented Reality
    left_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.4))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = c_card
    left_card.line.color.rgb = c_border
    left_card.line.width = Pt(1.5)

    ltf = left_card.text_frame
    ltf.word_wrap = True
    ltf.margin_left = ltf.margin_right = ltf.margin_top = Inches(0.4)

    lp0 = ltf.paragraphs[0]
    lp0.text = "CURRENT REALITY: FRAGMENTED & MANUAL"
    lp0.font.size = Pt(11)
    lp0.font.bold = True
    lp0.font.color.rgb = c_red

    lp1 = ltf.add_paragraph()
    lp1.text = "3 to 4 Separate Track Closures Per Section"
    lp1.font.size = Pt(16)
    lp1.font.bold = True
    lp1.font.color.rgb = c_navy
    lp1.space_before = Pt(4)
    lp1.space_after = Pt(12)

    l_bullets = [
        "• Civil Engineering requests 90 min for track tamping via BDMS.",
        "• 2 days later, TRD requests 120 min for OHE catenary power cut on the same stretch.",
        "• Signal & Telecom requests 60 min for point machine testing next week.",
        "• Result: High cumulative line downtime, severe freight train delays, and poor asset utilization."
    ]
    for b in l_bullets:
        p_b = ltf.add_paragraph()
        p_b.text = b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = c_slate
        p_b.space_after = Pt(8)

    # Right: Shadow Blocking
    right_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.933), Inches(1.5), Inches(5.6), Inches(5.4))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = c_card
    right_card.line.color.rgb = c_blue
    right_card.line.width = Pt(2)

    rtf = right_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_right = rtf.margin_top = Inches(0.4)

    rp0 = rtf.paragraphs[0]
    rp0.text = "OUR INNOVATION: CO-LOCATION & SHADOW BLOCKING"
    rp0.font.size = Pt(11)
    rp0.font.bold = True
    rp0.font.color.rgb = c_blue

    rp1 = rtf.add_paragraph()
    rp1.text = "1 Unified Multi-Department Mega Block"
    rp1.font.size = Pt(16)
    rp1.font.bold = True
    rp1.font.color.rgb = c_navy
    rp1.space_before = Pt(4)
    rp1.space_after = Pt(12)

    r_bullets = [
        "• Automatic Spatial Clustering: Cross-references chainage KM across TMS, SMMS, and TDMS.",
        "• Synchronized Execution: Bundles OHE power cut + Track tamping + Signal checks into 1 single closure.",
        "• COA Gap Alignment: Schedules the joint block precisely during natural off-peak freight gaps (01:30 - 03:30 AM).",
        "• Outcome: 38% reduction in corridor downtime with zero passenger disruptions."
    ]
    for b in r_bullets:
        p_b = rtf.add_paragraph()
        p_b.text = b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = c_slate
        p_b.space_after = Pt(8)

    # =========================================================================
    # SLIDE 5: Multi-Horizon Planning & Execution
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "Multi-Horizon Planning: Tactical 24h to Monthly Strategy")

    horizons = [
        ("Daily Tactical (24-Hour)", "Emergency & High Priority", [
            "• Immediate resolution of P1 critical safety defects (rail fractures, point failures).",
            "• Dynamic slot reallocation matching real-time train delay updates from COA.",
            "• Live safety buffer timers for field gangs and maintenance crews."
        ], c_blue),
        ("Weekly Operational (7-Day)", "Preventive Joint Bundling", [
            "• Multi-department task clustering across Track, OHE, and Signaling.",
            "• Pre-allocation of machine tamping and TRD tower wagons.",
            "• Conflict-free freight path planning with zero passenger timetable disturbance."
        ], c_dark_blue),
        ("Monthly Strategic (30-Day)", "Cyclical Overhaul & Asset Life", [
            "• Long-term corridor asset availability optimization.",
            "• Machine maintenance master schedule balancing division resource capacity.",
            "• RDSO regulatory compliance forecasting and executive audit reporting."
        ], c_green)
    ]

    for i, (title_h, sub_h, items_h, accent_h) in enumerate(horizons):
        cx = Inches(0.8 + i * 4.04)
        c_shape = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.5), Inches(3.64), Inches(5.4))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = c_card
        c_shape.line.color.rgb = c_border
        c_shape.line.width = Pt(1.5)

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)

        p_sub = tf.paragraphs[0]
        p_sub.text = sub_h.upper()
        p_sub.font.size = Pt(9)
        p_sub.font.bold = True
        p_sub.font.color.rgb = accent_h

        p_t = tf.add_paragraph()
        p_t.text = title_h
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = c_navy
        p_t.space_before = Pt(4)
        p_t.space_after = Pt(14)

        for itm in items_h:
            pi = tf.add_paragraph()
            pi.text = itm
            pi.font.size = Pt(11)
            pi.font.color.rgb = c_slate
            pi.space_after = Pt(10)

    # =========================================================================
    # SLIDE 6: Measurable Impact & Scalability
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "Operational Impact & Seamless Indian Railways Integration")

    impact_cards = [
        ("Asset Availability", "Up to +18% Increase", "Consolidating 3 separate blocks into 1 joint window keeps critical trunk corridors open for freight and passenger operations.", c_blue),
        ("Punctuality Index", "Zero Timetable Collisions", "Algorithm avoids passenger peak traffic and restricts major mega-blocks to low-density night slots.", c_navy),
        ("Controller Velocity", "< 30s Fast Solver", "Replaces multi-day inter-departmental meetings with an instantaneous constraint-satisfaction schedule.", c_green),
        ("Safety & Compliance", "100% Tamper-Evident", "Every approved block generates an immutable RDSO-compliant audit record with complete justification logs.", c_dark_blue)
    ]

    for i, (title_imp, stat_imp, desc_imp, acc_imp) in enumerate(impact_cards):
        row = i // 2
        col = i % 2
        cx = Inches(0.8 + col * 6.0)
        cy = Inches(1.5 + row * 2.8)

        card_imp = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(5.7), Inches(2.5))
        card_imp.fill.solid()
        card_imp.fill.fore_color.rgb = c_card
        card_imp.line.color.rgb = c_border
        card_imp.line.width = Pt(1.5)

        tf = card_imp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)

        p_t = tf.paragraphs[0]
        p_t.text = title_imp.upper()
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = acc_imp

        p_s = tf.add_paragraph()
        p_s.text = stat_imp
        p_s.font.size = Pt(18)
        p_s.font.bold = True
        p_s.font.color.rgb = c_navy
        p_s.space_before = Pt(2)
        p_s.space_after = Pt(6)

        p_d = tf.add_paragraph()
        p_d.text = desc_imp
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = c_slate

    # Save
    output_path = os.path.abspath("SIH_Automatic_Block_Planning_Presentation.pptx")
    prs.save(output_path)
    print(f"Complete 6-Slide Presentation generated at: {output_path}")

if __name__ == "__main__":
    create_complete_sih_deck()
